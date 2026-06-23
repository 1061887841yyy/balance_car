from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "outputs" / "FINAL_满版图文版_报告.docx"
IMG = ROOT / "outputs" / "图片材料"
ASSETS = ROOT / "outputs" / "food_robot_assets"
ASSETS.mkdir(exist_ok=True)

TITLE = "食品供应链双轮运输机器人设计与实现"


def load_font(size):
    for path in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"]:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def crop_cover(path, size):
    im = Image.open(path).convert("RGB")
    return ImageOps.fit(im, size, method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))


def make_plate(items, out_name, title, cell=(520, 300), label_h=38):
    font = load_font(24)
    title_font = load_font(28)
    pad = 18
    title_h = 42
    w = len(items) * cell[0] + (len(items) + 1) * pad
    h = title_h + cell[1] + label_h + 2 * pad
    sheet = Image.new("RGB", (w, h), (250, 250, 250))
    draw = ImageDraw.Draw(sheet)
    draw.text((pad, 8), title, fill=(25, 25, 25), font=title_font)
    draw.line((pad, title_h, w - pad, title_h), fill=(170, 170, 170), width=2)
    for i, (file_name, label) in enumerate(items):
        x = pad + i * (cell[0] + pad)
        y = title_h + pad
        draw.rectangle([x - 1, y - 1, x + cell[0] + 1, y + cell[1] + label_h + 1], outline=(185, 185, 185), width=2)
        im = crop_cover(IMG / file_name, cell)
        sheet.paste(im, (x, y))
        draw.text((x + cell[0] / 2, y + cell[1] + 8), label, fill=(20, 20, 20), font=font, anchor="ma")
    out = ASSETS / out_name
    sheet.save(out, quality=92)
    return out


hardware_plate = make_plate(
    [
        ("平衡小车废案.jpg", "双轮紧凑底盘"),
        ("主控stm32板.jpg", "STM32主控"),
        ("自研扩展版.png", "扩展板设计"),
    ],
    "plate_food_hardware.jpg",
    "食品运输机器人硬件平台",
)
sensor_plate = make_plate(
    [
        ("温湿度传感器.jpg", "温湿度环境监测"),
        ("柔性传感器.jpg", "食材重量检测"),
        ("oled.jpg", "OLED现场显示"),
    ],
    "plate_food_sensors.jpg",
    "质量与环境测量模块",
)
control_plate = make_plate(
    [
        ("蓝牙模块.jpg", "蓝牙通信模块"),
        ("蓝牙控制手机app.jpg", "手机APP远程控制"),
        ("模块支撑板.png", "模块支撑结构"),
    ],
    "plate_food_control.jpg",
    "远程控制与结构支撑",
)
future_plate = make_plate(
    [
        ("扩展版原理图1.png", "扩展板原理图"),
        ("扩展版原理图2.png", "接口与电源设计"),
        ("平衡小车废案.jpg", "导航运输构想"),
    ],
    "plate_food_future.jpg",
    "未来智能化扩展资料",
)


def set_run_font(run, east="宋体", size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def format_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, east="宋体", bold=False, first_line=True, line_pt=24, before=0, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(24) if first_line else None
    for run in p.runs:
        set_run_font(run, east, size, bold)


def clear_para(p):
    p.clear()


def new_para_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]


def remove_paragraph(p):
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def make_top_bottom_anchor(inline):
    anchor = inline
    anchor.tag = qn("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "251659264")
    anchor.set("behindDoc", "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "0")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    pos_offset = OxmlElement("wp:posOffset")
    pos_offset.text = "0"
    position_v.append(pos_offset)
    wrap = OxmlElement("wp:wrapTopAndBottom")

    children = list(anchor)
    for child in children:
        anchor.remove(child)
    for child in [simple_pos, position_h, position_v]:
        anchor.append(child)
    for tag in [qn("wp:extent"), qn("wp:effectExtent")]:
        found = next((c for c in children if c.tag == tag), None)
        if found is not None:
            anchor.append(found)
    anchor.append(wrap)
    for tag in [qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")]:
        found = next((c for c in children if c.tag == tag), None)
        if found is not None:
            anchor.append(found)


def add_wrapped_image_after(paragraph, image_path, caption, width_cm=11.6):
    p_img = new_para_after(paragraph)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_img.paragraph_format.line_spacing = Pt(12)
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(0)
    run = p_img.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    inline = run._element.xpath(".//wp:inline")[0]
    make_top_bottom_anchor(inline)
    p_cap = new_para_after(p_img)
    p_cap.add_run(caption)
    format_para(p_cap, WD_ALIGN_PARAGRAPH.CENTER, size=10.5, first_line=False, line_pt=16, before=0, after=4)
    return p_cap


def add_text_after(paragraph, text, role="body"):
    p = new_para_after(paragraph)
    p.add_run(text)
    if role == "heading":
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, size=14, east="楷体", bold=True, first_line=False, line_pt=24, before=8, after=2)
    elif role == "title":
        format_para(p, WD_ALIGN_PARAGRAPH.CENTER, size=16, east="黑体", bold=True, first_line=False, line_pt=24, before=0, after=4)
    else:
        format_para(p)
    return p


def find_para(doc, contains, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if contains in p.text:
            return i, p
    raise ValueError(f"not found: {contains}")


sections = [
    (
        "一、制作动机",
        [
            "本项目以食品供应链末端搬运场景为背景，设计一台食品供应链双轮运输机器人。传统食品搬运依赖人工手推车或固定货架转运，在走廊、厨房后场、仓储通道等狭窄空间中转弯半径较大，且无法实时反馈食材重量与环境状态。双轮底盘的优势在于结构紧凑、占用宽度小，能够在较小空间内完成转向和短距离运输，更适合餐饮供应、冷链中转、实验室样品配送等需要频繁移动的小型场景。",
            "食品供应链管理不仅关注“把物品送到”，还关注“送到的食材是否足称、运输环境是否合适”。因此，本项目在双轮底盘上加入温湿度传感器、重量称重传感器和OLED显示模块。温湿度数据用于判断当前环境是否适合食材短时间存放或运输，重量数据用于检查被搬运食材是否达到预期质量，避免缺斤少两或装载异常。OLED负责现场显示温湿度、重量和原始测量值，便于调试和展示。",
            "操控方式方面，本项目当前采用蓝牙APP远程控制。相比完全自动导航，蓝牙控制实现难度更适合本阶段作品展示，也能在食品供应链场景中完成基础的人机协同：操作者通过手机控制机器人前进、后退和转向，机器人负责承载食材并同步显示重量与环境状态。这样既降低了初版实现风险，也为后续加入自动导航、RFID识别和视觉识别留下接口。",
            "从课程学习角度看，本项目把传感与测量技术从单纯读数扩展到具体应用判断：温湿度不是孤立数据，而是食品保鲜环境依据；重量不是普通压力数值，而是食材足称检查依据；蓝牙控制不是单纯遥控，而是供应链搬运流程中的远程操作入口。项目目标是完成一台具备运输、测量、显示和远程控制能力的食品供应链双轮运输机器人原型。",
        ],
    ),
    (
        "二、文献探讨",
        [
            "食品运输与储存过程中的质量控制通常包含重量、温度、湿度、时间和环境暴露等因素。不同种类食材对保鲜环境要求不同，例如蔬菜、水果、肉类、乳制品和熟食在适宜温湿度范围、允许暴露时间和运输条件上存在差异。虽然本项目尚未建立完整食材数据库，但通过温湿度传感器采集环境数据，可以为后续判断“当前环境是否适合长时间放置该食品”提供基础测量入口。",
            "重量检测在食品供应链中具有验收与复核意义。食材进入运输环节前，常需要确认质量是否达标，避免缺斤少两或取放错误。本项目使用柔性压力传感器配合调理模块读取模拟量，再通过校准系数估算重量。该方案适合教学与原型验证，能够展示重量随放置食材变化而变化的趋势，但其精度受受力面积、安装方式和传感器非线性影响，后续若面向实际称重可更换为应变片称重传感器与HX711模块。",
            "温湿度测量采用DHT11模块。该模块通信简单、成本低，适合在课程项目中快速完成环境数据采集。它的精度和响应速度有限，因此更适合做环境状态提示，而不适合直接作为食品安全判定的唯一依据。在本项目中，DHT11的作用是提供当前运输环境的温湿度显示，并为未来保鲜环境判断和报警逻辑提供数据来源。",
            "双轮底盘结构常用于小型移动机器人，优点是占地面积小、转向灵活、结构简单。本项目采用双轮底盘并不是为了强调自平衡控制本身，而是为了在食品供应链末端搬运中减少空间占用。与四轮小车相比，双轮底盘更容易在狭窄通道内调整姿态；与固定输送线相比，它具有移动灵活和部署成本低的优点。",
            "蓝牙APP远程控制是当前阶段的人机交互方式。蓝牙模块与STM32主控通信后，手机端按键指令可以转换为电机控制命令，实现前进、后退、左转、右转和停止等基础动作。该方式虽然还不是完全自主导航，但能满足原型阶段的远程控制需求，并可作为后续自动导航失效时的人工接管方式。",
        ],
    ),
    (
        "三、设计内容",
        [
            "本作品主控芯片为STM32F103C8T6，系统围绕“食品运输、质量测量、环境显示、远程控制”四个功能展开。硬件上，双轮底盘承担移动运输任务，TB6612电机驱动模块控制左右轮运动；温湿度传感器采集当前环境状态；柔性压力传感器与调理模块输出模拟电压，经ADC读取后估算食材重量；OLED显示屏显示温度、湿度、重量和ADC值；蓝牙模块接收手机APP控制指令。",
            "双轮底盘设计的主要目的在于减少空间占用。食品供应链末端场景往往存在货架间距小、通道窄、转运路线短等特点，传统推车在这些环境中不够灵活。本项目使用左右轮差速控制，通过改变两侧电机转速实现前进、后退和转向，使机器人能够在有限空间内完成食材搬运动作。",
            "重量测量部分由柔性压力传感器和RFP602调理模块组成。食材放置在承载区域后，压力变化会反映为模拟输出电压，STM32通过PA2/ADC通道读取原始值，并根据空载零点和比例系数估算重量。OLED显示重量值和ADC原始值，便于在Ozone或现场调试中校准零点与系数。当前重量结果定位为估算值，用于判断是否接近目标质量。",
            "温湿度测量部分用于判断运输环境是否适合食材暂存。DHT11模块通过单总线方式输出温度和湿度，系统低频读取数据，避免影响电机控制和显示刷新。当前版本主要显示实时温湿度；未来可根据食材类别建立阈值表，例如对生鲜、熟食、蔬果分别设定适宜温湿度范围，超出范围时提示不宜长时间放置。",
            "蓝牙控制部分由蓝牙模块和手机APP组成。操作者可以通过手机端界面发送运动指令，主控收到指令后控制电机驱动模块执行对应动作。蓝牙远程控制让机器人能够在演示和实际搬运中由人进行调度，避免初版系统直接依赖复杂自动导航，也使作品更贴近“食品供应链末端人工辅助运输”的现实场景。",
            "显示与调试部分采用OLED和Ozone变量观察相结合。OLED面向现场展示，显示Temp、Humi、Weight、ADC等关键数据；Ozone面向调试，可观察传感器原始值、重量换算系数、故障标志和控制状态。通过现场显示和在线调试结合，可以快速判断是传感器接线、校准参数还是控制指令出现问题。",
        ],
    ),
    (
        "四、作品展示",
        [
            "作品上电后，操作者首先通过蓝牙APP连接机器人，确认控制指令能够正常下发。手机端控制前进、后退、左转、右转和停止，主控根据指令驱动双轮底盘运动。相比完全依赖线控或按键，蓝牙APP更符合食品运输机器人在实际使用中的远程调度需求，也便于展示时进行安全控制。",
            "食材质量检查流程如下：将待运输食材放置到承载区域，压力传感器产生电压变化，系统读取ADC原始值并换算为估算重量。若目标是检查食材是否足称，可先记录空载零点，再放置已知重量样品校准比例系数，之后通过OLED观察Weight数值是否接近目标质量。由于当前使用柔性压力传感器，结果主要用于趋势判断和课程展示。",
            "环境状态检查流程如下：DHT11周期读取温度与湿度，OLED显示当前环境数据。对于需要保鲜的食材，温湿度显示能够提醒操作者当前环境是否适合长时间放置。当前版本尚未写入不同食材的保鲜阈值，因此不会自动判定具体食材是否安全，但已经完成环境数据采集与显示入口。",
            "硬件展示中，主控板、蓝牙模块、温湿度模块、压力传感器、OLED和双轮底盘共同构成食品供应链双轮运输机器人原型。扩展板和模块支撑结构用于整理连接关系和固定模块，减少线路松动与临时搭建带来的不稳定性。整体作品从“能移动”扩展到“能测量、能显示、能远程控制”。",
            "调试时需要重点观察三类状态：第一是蓝牙连接和控制命令是否稳定；第二是温湿度数据是否周期更新；第三是压力传感器ADC值是否随食材放置发生变化。若OLED能显示温湿度和重量，蓝牙APP能控制底盘运动，说明食品供应链运输机器人原型的核心链路已经打通。",
        ],
    ),
    (
        "五、未来展望与心得",
        [
            "未来首先需要完善食品保鲜环境判断。当前系统只显示温湿度，尚未建立食材类别与适宜环境之间的对应关系。后续可建立食材保鲜数据库，将蔬菜、水果、肉类、乳制品、熟食等类别对应到不同温湿度范围和允许放置时间。当环境不合适时，机器人可通过OLED、蜂鸣器或APP提示报警，提醒操作者尽快处理。",
            "第二个方向是加入导航和RFID识别。食品供应链场景中，机器人不仅要移动，还要知道食材从哪里来、要送到哪里去。后续可在仓储或厨房区域布置RFID标签或导航标记，机器人读取标签后判断当前位置和目标点，实现从人工遥控向半自动运输过渡。当环境不适合长时间放置某类食品时，机器人可导航到更适合的存储地点。",
            "第三个方向是加入视觉识别系统。通过摄像头和YOLO模型识别食材种类后，系统可以自动匹配该食材的保鲜条件和重量验收规则。例如识别到蔬菜、水果或盒装食品后，系统自动调用对应温湿度阈值和目标重量范围。该功能目前尚未实现，适合作为后续智能化扩展方向写入废案与展望部分。",
            "第四个方向是提升称重精度和机械结构可靠性。柔性压力传感器适合做压力趋势检测，但若要用于实际供应链验收，应考虑使用称重传感器和HX711模块，并设计稳定的承载平台，使食材重量能够均匀作用在传感器上。同时，底盘需要增加防滑、限位和重心优化设计，避免运输过程中食材晃动影响测量。",
            "通过本项目，我们认识到传感与测量技术必须结合具体场景才有完整意义。温湿度、重量、蓝牙控制、底盘运动这些模块单独看都只是基础功能，但放到食品供应链运输机器人中，它们分别对应保鲜环境、足称验收、远程调度和空间适应能力。项目虽然仍是原型，但已经形成了面向食品供应链末端搬运的系统化设计思路。",
        ],
    ),
]


def delete_body_from(doc, start_index):
    for idx in range(len(doc.paragraphs) - 1, start_index, -1):
        remove_paragraph(doc.paragraphs[idx])


doc = Document(DOCX)

# Update front matter title while preserving group/member tables.
for p in doc.paragraphs:
    if p.text.startswith("题目："):
        clear_para(p)
        p.add_run(f"题目：{TITLE}")
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, size=12, first_line=False)
    if "【传感与测量技术期末项目制作】制作与报告" in p.text:
        clear_para(p)
        p.add_run("【食品供应链双轮运输机器人】制作与报告")
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, size=12, first_line=False)

start_idx, title_para = find_para(doc, "基于STM32F103C8T6")
delete_body_from(doc, start_idx)
clear_para(title_para)
title_para.add_run(TITLE)
format_para(title_para, WD_ALIGN_PARAGRAPH.CENTER, size=16, east="黑体", bold=True, first_line=False, line_pt=24, after=4)
current = title_para

for heading, paras in sections:
    current = add_text_after(current, heading, "heading")
    for para in paras:
        current = add_text_after(current, para, "body")
        if "本作品主控芯片为STM32F103C8T6" in para:
            current = add_wrapped_image_after(current, hardware_plate, "图1  食品供应链双轮运输机器人硬件平台", 11.6)
        if "重量测量部分由柔性压力传感器" in para:
            current = add_wrapped_image_after(current, sensor_plate, "图2  食材重量、温湿度与现场显示模块", 11.6)
        if "蓝牙控制部分由蓝牙模块" in para:
            current = add_wrapped_image_after(current, control_plate, "图3  蓝牙模块、手机APP控制与支撑结构", 11.6)
        if "第二个方向是加入导航和RFID识别" in para:
            current = add_wrapped_image_after(current, future_plate, "图4  RFID导航、视觉识别与保鲜判断的未来扩展构想", 11.6)

doc.save(DOCX)

with ZipFile(DOCX) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
    cell.text for table in doc.tables for row in table.rows for cell in row.cells
)
print(f"saved={DOCX}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"page_breaks={xml.count('w:type=\"page\"')} anchors={xml.count('<wp:anchor')} wrapTopAndBottom={xml.count('<wp:wrapTopAndBottom')}")
print(f"title_ok={TITLE in all_text}")
print(f"group_ok={'第05组' in all_text and '陈俊杰' in all_text and '冯超贤' in all_text and '徐浤彬' in all_text}")
print(f"old_theme_left={'多传感器自平衡小车' in all_text or '自平衡小车设计与实现' in all_text}")
print(f"future_ok={'YOLO' in all_text and 'RFID' in all_text and '保鲜' in all_text}")
print(f"qmarks={all_text.count('?')} placeholder={'此处写' in all_text}")
