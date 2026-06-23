from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs" / "FINAL_无问号_8页正文_报告.docx"
OUT = ROOT / "outputs" / "FINAL_图文增强_8页正文_报告.docx"
IMG = ROOT / "outputs" / "图片材料"


def set_run_font(run, east="宋体", size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def format_caption(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(2)
    pf.space_after = Pt(6)
    for run in p.runs:
        set_run_font(run, "宋体", 10.5)


def format_body(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(24)
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, "宋体", 12)


def new_para_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]


def add_image_after(paragraph, image_path, caption, width_cm=10.5):
    p_img = new_para_after(paragraph)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    p_cap = new_para_after(p_img)
    p_cap.add_run(caption)
    format_caption(p_cap)
    return p_cap


def find_para(doc, contains, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if contains in p.text:
            return i, p
    raise ValueError(f"not found: {contains}")


doc = Document(SRC)

# Hardware section: add real hardware photos after the hardware connection paragraph.
_, p = find_para(doc, "硬件连接方面")
anchor = add_image_after(p, IMG / "主控stm32板.jpg", "图1  STM32主控板与扩展连接实物", 9.8)
anchor = add_image_after(anchor, IMG / "mpu6050.jpg", "图2  MPU6050姿态传感器安装与接线", 9.8)
anchor = add_image_after(anchor, IMG / "自研扩展版.png", "图3  自研扩展板设计效果图", 10.5)

# Sensor/display section: add photos of DHT11, FSR and OLED.
_, p = find_para(doc, "DHT11驱动的设计重点")
anchor = add_image_after(p, IMG / "温湿度传感器.jpg", "图4  DHT11温湿度模块实物连接", 9.5)
anchor = add_image_after(anchor, IMG / "柔性传感器.jpg", "图5  柔性压力传感器与受力测试", 8.8)
anchor = add_image_after(anchor, IMG / "oled.jpg", "图6  OLED显示模块实物连接", 9.5)

# Work display section: add schematic and support board material.
_, p = find_para(doc, "作品上电后")
anchor = add_image_after(p, IMG / "扩展版原理图1.png", "图7  扩展板原理图设计资料", 11.5)
anchor = add_image_after(anchor, IMG / "模块支撑板.png", "图8  模块支撑板结构设计", 10.5)

# Future outlook: add abandoned balance navigation delivery concept and a short paragraph.
idx, p = find_para(doc, "后续改进可以从控制性能")
future_para = new_para_after(p)
future_para.add_run(
    "    在作品构想阶段，团队还曾规划过“平衡导航食品运输小车”方案：在自平衡底盘上增加雷达或循迹模块，使小车能够沿指定路径完成食品运输与供应。由于展示时间、调试风险和硬件稳定性限制，最终没有把该方案作为本次展示重点，但它为后续扩展提供了明确方向。未来可在现有姿态控制和传感器显示基础上，继续加入路径识别、避障、载物平台和通信上传功能，使作品从测量演示平台进一步发展为具备场景任务能力的移动服务平台。"
)
format_body(future_para)
add_image_after(future_para, IMG / "平衡小车废案.jpg", "图9  废案构想：平衡导航食品运输小车", 7.4)

doc.save(OUT)

with ZipFile(OUT) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
text = "\n".join(p.text for p in doc.paragraphs)
print(f"saved={OUT}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"page_breaks={xml.count('w:type=\"page\"')} sectPr={xml.count('<w:sectPr')}")
print(f"qmarks={text.count('?')} placeholder={'此处写' in text}")
