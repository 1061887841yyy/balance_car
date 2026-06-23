import zipfile

from docx import Document


p = "outputs/FINAL_满版图文版_报告.docx"
d = Document(p)
txt = "\n".join(x.text for x in d.paragraphs)
with zipfile.ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8")

print("paras", len(d.paragraphs))
print("imgs", len(d.inline_shapes))
print("tables", len(d.tables))
print("qmarks", txt.count("?"))
print("placeholder", "此处写" in txt)
print("pageBreaks", xml.count('w:type="page"'))
print("sectPr", xml.count("<w:sectPr"))
print("future", "平衡导航食品运输小车" in txt)
