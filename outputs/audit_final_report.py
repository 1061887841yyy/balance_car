from zipfile import ZipFile

from docx import Document


path = "outputs/FINAL_满版图文版_报告.docx"
doc = Document(path)
para_text = "\n".join(p.text for p in doc.paragraphs)
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
all_text = para_text + "\n" + table_text

with ZipFile(path) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")

print("paragraphs", len(doc.paragraphs))
print("tables", len(doc.tables))
print("inline_shapes", len(doc.inline_shapes))
print("page_breaks", xml.count('w:type="page"'))
print("anchors", xml.count("<wp:anchor"))
print("wrapTopAndBottom", xml.count("<wp:wrapTopAndBottom"))
print("group05", "第05组" in all_text)
print("members", all(name in all_text for name in ["陈俊杰", "冯超贤", "徐浤彬"]))
print("ids", all(sid in all_text for sid in ["2024433080104", "2024433080110", "2024433080138"]))
print("bad_meta_sentence", "减少单张照片独占页面" in all_text or "这样的图文组合" in all_text)
print("future_caption", "图4  废案构想：平衡导航食品运输小车" in all_text)
print("placeholder", "此处写" in all_text)
print("qmarks", all_text.count("?"))
