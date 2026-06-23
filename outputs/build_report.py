from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "source_template.docx"
OUT = ROOT / "outputs" / "report_filled_fixed.docx"
OUT_CN = ROOT / "outputs" / "《传感与测量技术》Z1期末报告_第XX组_已填写.docx"


TITLE = "基于STM32F103C8T6的多传感器自平衡小车设计与实现"


PAGES = [
    {
        "heading": "一、制作动机",
        "paragraphs": [
            "本项目选择两轮自平衡小车作为期末作品，是因为它能够把传感与测量技术、嵌入式控制、电机驱动和调试校准过程集中到一个可观察、可验证的系统中。普通的传感器实验往往只停留在读数显示，而自平衡小车要求传感器数据必须直接参与控制闭环，任何姿态角、速度值或采样周期的误差都会立刻反映到车体晃动、失稳或电机输出异常上，因此更能体现测量数据在工程系统中的实际价值。",
            "在原有平衡车基础上，本次设计进一步加入DHT11温湿度传感器、FSR402压力传感器配合RFP602调理模块，以及I2C接口OLED显示屏。这样，小车不仅能够依靠MPU6050和编码器完成姿态与速度测量，还能够采集周围环境温湿度，并估算放置在压力传感器上的物体重量。OLED用于本地显示温度、湿度、重量和ADC原始值，Ozone与J-Link用于在线观察全局变量和调节参数，使作品同时具备运动控制、环境测量、压力测量和人机显示功能。",
            "从学习角度看，本作品的制作动机主要有三点。第一，通过MPU6050姿态采集、编码器测速和TIM定时节拍，理解动态测量系统中采样频率、滤波方法和执行机构响应之间的关系。第二，通过DHT11和FSR402的加入，比较数字单总线传感器与模拟量传感器在读取方式、误差来源和校准方法上的差别。第三，通过OLED显示和Ozone变量观察，把实验结果从隐藏在程序内部的数据变成可以现场检查的状态量，从而提高调试效率和作品展示效果。",
            "本项目最终希望实现一个可运行、可调试、可扩展的STM32多传感器平台。它既可以作为自平衡控制实验平台，也可以作为温湿度采集、压力估算和显示交互的综合案例。相比单独完成一个传感器读数实验，这种组合式作品更接近真实工程系统，因为真实产品通常需要多个传感器并行工作，并在有限的处理器资源和有限的时间预算下完成采样、计算、显示和控制。",
        ],
    },
    {
        "heading": "二、文献探讨",
        "paragraphs": [
            "两轮自平衡小车的理论模型可近似看作倒立摆系统。倒立摆本身是不稳定对象，当车体向前或向后倾斜时，控制器需要驱动轮子朝倾倒方向运动，使支撑点追上车体重心，从而维持动态平衡。该类系统通常需要快速、连续的姿态测量和电机闭环控制，因此陀螺仪、加速度计和编码器是平衡车设计中最常见的核心传感器。",
            "MPU6050内部集成三轴加速度计和三轴陀螺仪。加速度计可以根据重力方向计算静态倾角，但在车体加减速或受到震动时会混入线加速度，瞬时读数容易抖动；陀螺仪可以测得角速度，积分后可得到角度变化，短时间响应平滑，但长时间会产生零偏漂移。工程中常采用互补滤波方法，将加速度计的长期稳定性与陀螺仪的短期动态响应结合起来。本项目中俯仰角采用加速度角与陀螺仪积分角按比例融合，得到用于角度环PID的姿态角。",
            "编码器测速用于速度环与转向环。直流电机仅靠PWM开环输出时，实际速度会受到电池电压、地面摩擦、负载和电机差异影响。通过TIM编码器模式读取左右轮脉冲增量，可以计算左右轮速度、平均速度和差速。平均速度反馈用于抑制小车长期向前或向后漂移，差速反馈用于控制左右轮一致性，为转向控制或直行稳定提供依据。",
            "DHT11是一种常见的温湿度数字传感器，采用单总线通信方式。主机先发送起始信号，随后DHT11返回响应并输出40位数据。该传感器成本低、连接简单，适合教学实验和普通环境监测，但采样速度较低，精度也有限。因此程序中采用低频读取策略，每约2秒更新一次温湿度数据，避免它占用平衡控制所需的高频计算时间。",
        ],
    },
    {
        "heading": "二、文献探讨（续）",
        "paragraphs": [
            "FSR402是一类薄膜压力敏感电阻，其阻值会随受力变化而变化。它适合检测按压强弱、接触状态或进行粗略重量估算，但其输出不是严格线性的，且受受力面积、材料回弹、安装方式和温度等因素影响较大。RFP602模块可将FSR402的阻值变化转换为模拟电压输出，STM32通过ADC1_IN2采集该电压，再换算为ADC原始值和毫伏值。",
            "压力传感器用于“重量g”显示时必须进行校准。项目采用的换算公式为weight_g=max(0,(fsr_adc_raw-fsr_zero_adc)*fsr_g_per_count)。其中fsr_zero_adc为空载时的零点ADC值，fsr_g_per_count为每个ADC计数对应的估算克重。该公式简单、便于在Ozone中观察和修改，适合教学演示，但并不等同于精密电子秤。实际使用时需要先空载记录零点，再放置已知重量物体，根据显示偏差调整比例系数。",
            "OLED显示屏采用SSD1306控制器，分辨率为128x64，I2C地址默认为0x3C。OLED的优点是体积小、功耗低、对比度高，适合显示少量调试信息。本项目将OLED与MPU6050分开使用不同I2C总线，MPU6050保留在I2C1的PB8/PB9，OLED改用I2C2的PB10/PB11，避免两个模块竞争同一组引脚，也便于排查硬件连接问题。",
            "综合相关资料可以看出，本项目的关键并不是单个传感器读数，而是多传感器数据在统一控制框架中的组织方式。高频姿态与编码器数据服务于平衡控制，低频温湿度和重量估算服务于显示与测量扩展。合理的节拍安排、故障标志和调试变量能够让系统在新增模块出现异常时仍尽量保持主控制逻辑稳定。",
        ],
    },
    {
        "heading": "三、设计内容",
        "paragraphs": [
            "本作品主控芯片为STM32F103C8T6，工程由STM32CubeMX生成Makefile项目后继续开发，核心代码位于Core/Inc/balance_car和Core/Src/balance_car目录。主程序在初始化阶段调用BalanceCar_Init()，在while循环中持续调用BalanceCar_Background()。定时器TIM4提供1ms系统节拍，后台任务根据节拍执行10ms角度环、50ms速度环和转向环，以及较低频率的传感器显示任务。",
            "硬件连接方面，MPU6050使用I2C1，SCL接PB8，SDA接PB9，地址为0x68。OLED为0.96寸SSD1306 128x64显示屏，改用I2C2，SCL接PB10，SDA接PB11，默认地址为0x3C。DHT11数据脚使用PC14，电源接3.3V并与系统共地。FSR402通过RFP602模块输出模拟电压，AO接PA2，即ADC1_IN2。TB6612电机驱动中，STBY接PA3，AIN1接PA4，AIN2接PA5，PWMA接PA6/TIM3_CH1，BIN1接PB1，BIN2接PB0，PWMB接PA7/TIM3_CH2。左编码器使用PA8/PA9对应TIM1，右编码器使用PA0/PA1对应TIM2。",
            "软件模块按照功能进行划分。balance_control.c负责系统初始化、控制节拍、调试变量和主控制逻辑；pid.c实现位置式PID，并提供积分限幅、输出限幅和死区补偿；mpu6050_hal.c通过HAL I2C读取加速度计和陀螺仪原始数据；encoder_hal.c使用定时器编码器模式读取轮速；motor_tb6612.c负责电机方向和PWM输出；i2c_bus.c统一初始化I2C1和I2C2；oled_ssd1306.c实现OLED命令、显存和分页刷新；dht11.c实现PC14单总线温湿度读取；fsr_adc.c负责PA2模拟量采集；app_sensors.c汇总温湿度、电压、重量估算和故障标志；display_ui.c负责OLED四行数据显示。",
            "系统初始化采用分层容错策略。BalanceCar_Init()首先初始化I2C、MPU6050、电机、编码器和PID相关参数，再初始化OLED、DHT11和FSR ADC等扩展模块。若OLED或某个传感器初始化失败，程序只设置sensor_fault_flags或对应ready变量，不直接阻断平衡控制初始化。这样在调试阶段即使显示屏未亮或DHT11接线错误，也可以继续通过Ozone观察变量并验证核心控制。",
        ],
    },
    {
        "heading": "三、设计内容（续）",
        "paragraphs": [
            "平衡控制采用串级思路。10ms角度环从MPU6050读取原始数据，计算加速度角angle_acc和陀螺仪积分角angle_gyro，再通过互补滤波得到当前俯仰角angle。角度PID以目标角度和实际角度为输入，输出平均PWM。50ms速度环从左右编码器读取脉冲差值，换算得到left_speed和right_speed，再计算ave_speed与dif_speed。速度PID输出作为角度环目标值，用于抑制小车长期漂移；转向PID根据左右速度差输出差分PWM，使左右轮运动更一致。",
            "新增传感器任务放在BalanceCar_Background()中执行，但不改变原有10ms和50ms控制节拍。FSR ADC约每100ms读取一次，读取后更新fsr_adc_raw和fsr_voltage_mv，并根据fsr_zero_adc与fsr_g_per_count计算weight_g。DHT11约每2秒读取一次，读取失败时保留上次有效温湿度并设置故障标志。OLED约每500ms重新组织显示文字，但实际刷新采用分页方式，每次只刷新一页或少量页面，减少I2C传输对平衡控制的阻塞。",
            "OLED默认显示四行数据：第一行为Temp: xx.x C，第二行为Humi: xx %，第三行为Weight: xxxx g，第四行为ADC: xxxx。温湿度来自DHT11，重量来自FSR估算，ADC原始值用于Ozone校准。保留ADC值是因为FSR402本身误差较大，单看重量数字难以判断是否为比例系数错误、零点漂移或传感器安装问题。通过同时观察ADC、毫伏值和重量，能更快定位问题。",
            "为了便于调试，程序提供了一组全局状态变量供Ozone观察，包括temperature_c、humidity_percent、fsr_adc_raw、fsr_voltage_mv、weight_g、fsr_zero_adc、fsr_g_per_count和sensor_fault_flags。OLED状态还可以通过oled_ready、oled_last_error、oled_i2c_addr等变量检查。若OLED不亮，仍可在Ozone中查看这些变量判断是显示屏初始化失败、I2C地址不匹配，还是传感器数据本身没有更新。",
        ],
    },
    {
        "heading": "三、设计内容（续）",
        "paragraphs": [
            "DHT11驱动的设计重点是避免长时间阻塞。DHT11通信需要主机拉低数据线约18ms作为起始信号，如果直接在主循环中延时等待，会影响控制任务执行。因此程序将低电平起始阶段设计为状态化过程，真正读取40位数据时再使用带超时的短等待。这样即使DHT11没有响应，也能在超时后退出，不会让主循环长时间停在一个传感器函数里。",
            "FSR ADC驱动使用ADC1单通道模式读取PA2。ADC采样后，程序按3.3V参考电压将12位ADC值换算为毫伏值，即fsr_voltage_mv=fsr_adc_raw*3300/4095。空载时程序记录或允许用户设置fsr_zero_adc，放置已知重量后通过Ozone修改fsr_g_per_count，使OLED显示值接近实物重量。由于FSR402受力不均会造成明显误差，实际测试时应尽量保证受力点稳定，并在同一安装方式下完成校准。",
            "I2C总线设计也进行了调整。最初OLED可能与MPU6050共用PB8/PB9，但用户要求OLED不要占用MPU6050使用的引脚，因此工程改为I2C1服务MPU6050，I2C2服务OLED。这样不仅硬件接线更清晰，也减少了显示刷新与姿态读取在同一总线上的竞争。I2C2引脚PB10/PB11需要在CubeMX或底层初始化中正确配置为复用开漏，并保证外部上拉或模块自带上拉正常。",
            "整体程序的设计原则是“控制优先、测量可见、故障可查”。控制优先指姿态和电机闭环的节拍不因显示和低速传感器被破坏；测量可见指关键传感器数据不仅在OLED显示，也保留在Ozone变量中；故障可查指初始化失败、读取失败、OLED错误和传感器异常都通过状态变量暴露出来，便于逐步排查硬件接线和软件配置。",
        ],
    },
    {
        "heading": "四、作品展示",
        "paragraphs": [
            "作品上电后，首先完成主控、定时器、电机驱动、编码器、MPU6050、OLED、DHT11和ADC模块初始化。正常情况下，OLED屏幕显示四行信息：Temp用于显示当前温度，Humi用于显示当前湿度，Weight用于显示估算重量，ADC用于显示压力传感器原始采样值。若温湿度模块连接正常，Ozone中temperature_c和humidity_percent约每2秒更新一次；若按压FSR402或在其上放置物体，fsr_adc_raw、fsr_voltage_mv和weight_g会随压力变化而改变。",
            "调试时可以先只连接OLED，确认oled_ready为1，oled_i2c_addr为0x3C，并观察屏幕是否能显示占位数值。如果OLED不亮，应检查PB10/PB11接线、电源、GND、I2C地址和屏幕方向。随后连接DHT11，观察temperature_c和humidity_percent是否更新，若sensor_fault_flags出现对应故障位，则检查PC14数据线和上拉电阻。最后连接RFP602输出到PA2，观察fsr_adc_raw是否会随按压增加或减小。",
            "压力传感器校准过程可分为两步。第一步为空载校准，将传感器上方不放物体，记录稳定后的fsr_adc_raw并写入fsr_zero_adc。第二步为比例校准，放置一个已知重量物体，例如100g或200g砝码，观察OLED显示重量，如果显示偏小则增大fsr_g_per_count，如果显示偏大则减小该系数。校准完成后，OLED上的Weight值可以作为估算重量显示，ADC值仍保留给进一步修正使用。",
            "在平衡控制演示中，可以通过Ozone同时观察angle、left_speed、right_speed、ave_speed、g_angle_pid.Out、g_speed_pid.Out和timer_error_flag等变量。若timer_error_flag不置位，说明新增OLED、DHT11和ADC任务没有造成明显节拍堆积；若出现控制不稳，应先断开扩展测量任务或降低显示刷新频率，再确认PID参数、电机方向、编码器方向和MPU6050安装方向是否正确。",
        ],
    },
    {
        "heading": "五、未来展望与心得",
        "paragraphs": [
            "后续改进可以从控制性能、测量精度和交互方式三个方向展开。控制性能方面，可以继续细调角度环、速度环和转向环PID参数，增加低电压保护、跌倒保护和软启动策略，使小车在不同地面和不同电池电压下更稳定。测量精度方面，FSR402适合压力趋势检测，但并不是精密称重元件。如果希望获得更可靠的重量值，可以改用应变片称重传感器和HX711放大采集模块，并进行多点标定。",
            "显示与通信方面，OLED目前只显示四行基础数据，后续可以增加按键切换页面，例如控制状态页、传感器状态页、PID参数页和故障诊断页。也可以加入串口、蓝牙或Wi-Fi模块，把温湿度、压力、姿态角和速度数据发送到上位机或手机端，形成更完整的数据记录系统。若需要进一步扩展作品功能，还可以增加循迹、避障或遥控模块，使平衡车从单一控制平台变成可移动测量平台。",
            "通过本项目的制作，我们体会到传感与测量并不是简单地把传感器接到单片机上读取数字。一个可用的测量系统需要考虑供电、接地、引脚复用、采样周期、滤波、校准、误差来源和数据显示方式。尤其在平衡车这类动态系统中，传感器数据会直接影响执行机构输出，因此任何延时、噪声或方向错误都会被放大成明显的运动问题。",
            "本次设计也让我们认识到模块化程序结构的重要性。将MPU6050、编码器、电机、OLED、DHT11、FSR ADC和显示界面分别封装后，调试时可以逐个验证模块，再将它们组合到主控制流程中。Ozone变量观察为调试提供了很大帮助，使我们能够在不反复修改显示代码的情况下查看内部状态。总体而言，本作品完成了从传感器采集、数据处理、控制输出到本地显示的完整链路，达到了课程综合实践的目标。",
        ],
    },
]


def set_rfonts(run, east_asia: str, size_pt: float, bold: bool = False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size_pt)
    run.bold = bold


def format_paragraph(paragraph, align=None, first_line=False):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_line:
        pf.first_line_indent = Pt(24)
    else:
        pf.first_line_indent = None
    paragraph.alignment = align


def clear_paragraph(paragraph):
    paragraph.clear()
    format_paragraph(paragraph)


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.add_paragraph()._p


def paragraph_after(paragraph):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]


def add_after(paragraph, text="", role="body"):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]
    write_para(new_para, text, role)
    return new_para


def write_para(paragraph, text, role):
    paragraph.clear()
    if role == "title":
        format_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run(text)
        set_rfonts(run, "黑体", 16, True)
    elif role == "heading":
        format_paragraph(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
        run = paragraph.add_run(text)
        set_rfonts(run, "楷体", 14, True)
    elif role == "body":
        format_paragraph(paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True)
        run = paragraph.add_run(text)
        set_rfonts(run, "宋体", 12, False)
    elif role == "blank":
        format_paragraph(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
    else:
        raise ValueError(role)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def page_break_after(paragraph):
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def build():
    doc = Document(SOURCE)

    # Remove the old placeholder paragraphs for the report body only. All front
    # matter, tables, section properties and earlier content are kept intact.
    for idx in range(62, 48, -1):
        remove_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[48]
    write_para(anchor, TITLE, "title")
    current = anchor

    for page_index, page in enumerate(PAGES):
        current = add_after(current, "", "blank")
        current = add_after(current, page["heading"], "heading")
        for text in page["paragraphs"]:
            current = add_after(current, text, "body")
        if page_index != len(PAGES) - 1:
            page_break_after(current)

    # Keep trailing blank template paragraphs, but make sure they have normal spacing.
    for p in doc.paragraphs:
        if not p.text.strip():
            format_paragraph(p)

    doc.save(OUT)
    doc.save(OUT_CN)

    with ZipFile(OUT) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    page_breaks = xml.count('<w:br w:type="page"')
    sect_pr = xml.count("<w:sectPr")
    placeholder_left = "此处写" in "\n".join(p.text for p in doc.paragraphs)
    print(f"saved={OUT}")
    print(f"saved_cn={OUT_CN}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
    print(f"page_breaks={page_breaks}")
    print(f"sectPr={sect_pr}")
    print(f"placeholder_left={placeholder_left}")


if __name__ == "__main__":
    build()
