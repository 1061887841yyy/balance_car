from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs" / "FINAL_无问号_8页正文_报告.docx"
OUT = ROOT / "outputs" / "FINAL_紧凑图文版_报告.docx"
IMG = ROOT / "outputs" / "图片材料"
WORK = ROOT / "outputs" / "report_compact_assets"
WORK.mkdir(exist_ok=True)


def load_font(size=28):
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def crop_cover(path, size):
    im = Image.open(path).convert("RGB")
    return ImageOps.fit(im, size, method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))


def make_strip(items, out_name, cell=(520, 340), label_h=44):
    font = load_font(24)
    pad = 22
    w = len(items) * cell[0] + (len(items) + 1) * pad
    h = cell[1] + label_h + 2 * pad
    sheet = Image.new("RGB", (w, h), (248, 248, 248))
    draw = ImageDraw.Draw(sheet)
    for i, (file_name, label) in enumerate(items):
        x = pad + i * (cell[0] + pad)
        y = pad
        draw.rectangle([x - 2, y - 2, x + cell[0] + 2, y + cell[1] + label_h + 2], outline=(190, 190, 190), width=2)
        im = crop_cover(IMG / file_name, cell)
        sheet.paste(im, (x, y))
        draw.text((x + cell[0] / 2, y + cell[1] + 10), label, fill=(20, 20, 20), font=font, anchor="ma")
    out = WORK / out_name
    sheet.save(out, quality=92)
    return out


hardware_strip = make_strip(
    [
        ("主控stm32板.jpg", "STM32主控板"),
        ("mpu6050.jpg", "MPU6050姿态模块"),
        ("自研扩展版.png", "自研扩展板"),
    ],
    "hardware_strip.jpg",
)

sensor_strip = make_strip(
    [
        ("温湿度传感器.jpg", "DHT11温湿度"),
        ("柔性传感器.jpg", "柔性压力传感器"),
        ("oled.jpg", "OLED显示屏"),
    ],
    "sensor_strip.jpg",
)

design_strip = make_strip(
    [
        ("扩展版原理图1.png", "扩展板原理图"),
        ("模块支撑板.png", "模块支撑板"),
        ("平衡小车废案.jpg", "导航运输废案"),
    ],
    "design_future_strip.jpg",
)


def set_run_font(run, east="宋体", size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def format_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, first_line=True, line_pt=24, before=0, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    pf.first_line_indent = Pt(24) if first_line else None
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    for run in p.runs:
        set_run_font(run, "宋体", size)


def new_para_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]


def add_compact_image_after(paragraph, image_path, caption, width_cm=13.2):
    p_img = new_para_after(paragraph)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(0)
    run = p_img.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    p_cap = new_para_after(p_img)
    p_cap.add_run(caption)
    format_para(p_cap, WD_ALIGN_PARAGRAPH.CENTER, size=10.5, first_line=False, line_pt=16, before=0, after=4)
    return p_cap


def find_para(doc, contains, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if contains in p.text:
            return i, p
    raise ValueError(f"not found: {contains}")


doc = Document(SRC)

_, p = find_para(doc, "硬件连接方面")
add_compact_image_after(p, hardware_strip, "图1  主控、姿态传感器与扩展板实物/设计资料", 12.2)

_, p = find_para(doc, "DHT11驱动的设计重点")
add_compact_image_after(p, sensor_strip, "图2  温湿度、压力与OLED显示模块实物", 12.2)

_, p = find_para(doc, "作品上电后")
add_compact_image_after(p, design_strip, "图3  扩展板设计资料与导航运输小车废案", 12.2)

_, p = find_para(doc, "后续改进可以从控制性能")
future_para = new_para_after(p)
future_para.add_run(
    "    在作品构想阶段，团队还曾规划过“平衡导航食品运输小车”方案：在自平衡底盘上增加雷达或循迹模块，使小车能够沿指定路径完成食品运输与供应。由于展示时间、调试风险和硬件稳定性限制，最终没有把该方案作为本次展示重点，但该废案为后续扩展提供了明确方向。"
)
format_para(future_para)

doc.save(OUT)

with ZipFile(OUT) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
text = "\n".join(p.text for p in doc.paragraphs)
print(f"saved={OUT}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"page_breaks={xml.count('w:type=\"page\"')} sectPr={xml.count('<w:sectPr')}")
print(f"qmarks={text.count('?')} placeholder={'此处写' in text} future={'平衡导航食品运输小车' in text}")
