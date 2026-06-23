from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "outputs" / "FINAL_无问号_8页正文_报告.docx"
OUT = ROOT / "outputs" / "FINAL_满版图文版_报告.docx"
IMG = ROOT / "outputs" / "图片材料"
ASSETS = ROOT / "outputs" / "report_full_assets"
ASSETS.mkdir(exist_ok=True)


def load_font(size):
    for path in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"]:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def crop_cover(path, size):
    im = Image.open(path).convert("RGB")
    return ImageOps.fit(im, size, method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))


def make_plate(items, out_name, title, cell=(520, 300), label_h=38):
    font = load_font(24)
    title_font = load_font(28)
    pad = 18
    title_h = 42
    w = len(items) * cell[0] + (len(items) + 1) * pad
    h = title_h + cell[1] + label_h + 2 * pad
    sheet = Image.new("RGB", (w, h), (250, 250, 250))
    draw = ImageDraw.Draw(sheet)
    draw.text((pad, 8), title, fill=(25, 25, 25), font=title_font)
    draw.line((pad, title_h, w - pad, title_h), fill=(170, 170, 170), width=2)
    for i, (file_name, label) in enumerate(items):
        x = pad + i * (cell[0] + pad)
        y = title_h + pad
        draw.rectangle([x - 1, y - 1, x + cell[0] + 1, y + cell[1] + label_h + 1], outline=(185, 185, 185), width=2)
        im = crop_cover(IMG / file_name, cell)
        sheet.paste(im, (x, y))
        draw.text((x + cell[0] / 2, y + cell[1] + 8), label, fill=(20, 20, 20), font=font, anchor="ma")
    out = ASSETS / out_name
    sheet.save(out, quality=92)
    return out


hardware_plate = make_plate(
    [
        ("主控stm32板.jpg", "STM32主控板"),
        ("mpu6050.jpg", "MPU6050姿态模块"),
        ("自研扩展版.png", "自研扩展板"),
    ],
    "plate_hardware_filled_base.jpg",
    "硬件实物与扩展板",
)

sensor_plate = make_plate(
    [
        ("温湿度传感器.jpg", "DHT11温湿度"),
        ("柔性传感器.jpg", "柔性压力传感器"),
        ("oled.jpg", "OLED显示屏"),
    ],
    "plate_sensor_filled_base.jpg",
    "温湿度、压力与显示模块",
)

design_plate = make_plate(
    [
        ("扩展版原理图1.png", "扩展板原理图"),
        ("模块支撑板.png", "模块支撑板"),
    ],
    "plate_design_filled_base.jpg",
    "扩展板与支撑板设计资料",
    cell=(700, 320),
)


def set_run_font(run, east="宋体", size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def format_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, first_line=True, line_pt=24, before=0, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(24) if first_line else None
    for run in p.runs:
        set_run_font(run, "宋体", size)


def new_para_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[[p._p for p in paragraph._parent.paragraphs].index(new_p)]


def make_top_bottom_anchor(inline):
    anchor = inline
    anchor.tag = qn("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "251659264")
    anchor.set("behindDoc", "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "0")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    pos_offset = OxmlElement("wp:posOffset")
    pos_offset.text = "0"
    position_v.append(pos_offset)

    wrap = OxmlElement("wp:wrapTopAndBottom")

    children = list(anchor)
    for child in children:
        anchor.remove(child)

    extent = next((c for c in children if c.tag == qn("wp:extent")), None)
    effect = next((c for c in children if c.tag == qn("wp:effectExtent")), None)
    doc_pr = next((c for c in children if c.tag == qn("wp:docPr")), None)
    cnv = next((c for c in children if c.tag == qn("wp:cNvGraphicFramePr")), None)
    graphic = next((c for c in children if c.tag == qn("a:graphic")), None)

    anchor.append(simple_pos)
    anchor.append(position_h)
    anchor.append(position_v)
    if extent is not None:
        anchor.append(extent)
    if effect is not None:
        anchor.append(effect)
    anchor.append(wrap)
    for child in [doc_pr, cnv, graphic]:
        if child is not None:
            anchor.append(child)


def add_wrapped_image_after(paragraph, image_path, caption, width_cm=12.0):
    p_img = new_para_after(paragraph)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_img.paragraph_format.line_spacing = Pt(12)
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(0)
    run = p_img.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    inline = run._element.xpath(".//wp:inline")[0]
    make_top_bottom_anchor(inline)

    p_cap = new_para_after(p_img)
    p_cap.add_run(caption)
    format_para(p_cap, WD_ALIGN_PARAGRAPH.CENTER, size=10.5, first_line=False, line_pt=16, before=0, after=4)
    return p_cap


def add_body_after(paragraph, text):
    p = new_para_after(paragraph)
    p.add_run(text)
    format_para(p)
    return p


def find_para(doc, contains, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if contains in p.text:
            return i, p
    raise ValueError(f"not found: {contains}")


def remove_body_page_breaks(doc):
    body_start, _ = find_para(doc, "基于STM32F103C8T6")
    for p in doc.paragraphs[body_start:]:
        for br in p._element.xpath(".//w:br[@w:type='page']"):
            br.getparent().remove(br)


doc = Document(BASE)
remove_body_page_breaks(doc)

_, p = find_para(doc, "硬件连接方面")
anchor = add_wrapped_image_after(p, hardware_plate, "图1  主控、姿态传感器与扩展板集中展示", 11.6)
anchor = add_body_after(
    anchor,
    "    图1展示了主控板、MPU6050姿态模块和自研扩展板的对应关系。主控板负责控制与数据处理，MPU6050负责姿态角测量，扩展板用于整理模块接口和后续扩展连接。"
)

_, p = find_para(doc, "DHT11驱动的设计重点")
anchor = add_wrapped_image_after(p, sensor_plate, "图2  温湿度、压力与OLED显示模块实物", 11.6)
anchor = add_body_after(
    anchor,
    "    图2对应本项目新增的测量与显示模块。DHT11用于温湿度采集，柔性压力传感器配合RFP602进行压力/重量估算，OLED用于现场显示Temp、Humi、Weight和ADC等关键数据。"
)

_, p = find_para(doc, "作品上电后")
anchor = add_wrapped_image_after(p, design_plate, "图3  扩展板原理图与模块支撑板设计资料", 11.6)
anchor = add_body_after(
    anchor,
    "    图3展示了扩展板原理图和模块支撑板设计资料。原理图用于说明电机驱动、传感器接口和电源输入等电路连接关系，支撑板用于固定温湿度、压力、OLED等扩展模块，使硬件布局更加清晰。"
)

_, p = find_para(doc, "后续改进可以从控制性能")
anchor = add_body_after(
    p,
    "    在作品构想阶段，团队还曾规划过“平衡导航食品运输小车”方案：在自平衡底盘上增加雷达或循迹模块，使小车能够沿指定路径完成食品运输与供应。由于展示时间、调试风险和硬件稳定性限制，最终没有把该方案作为本次展示重点，但它为后续扩展提供了明确方向。"
)
add_wrapped_image_after(anchor, IMG / "平衡小车废案.jpg", "图4  废案构想：平衡导航食品运输小车", 6.0)

doc.save(OUT)

with ZipFile(OUT) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
text = "\n".join(p.text for p in doc.paragraphs)
print(f"saved={OUT}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"page_breaks={xml.count('w:type=\"page\"')} anchors={xml.count('<wp:anchor')} wrapTopAndBottom={xml.count('<wp:wrapTopAndBottom')}")
print(f"group_info={'陈俊杰' in text and '冯超贤' in text and '徐浤彬' in text}")
print(f"qmarks={text.count('?')} placeholder={'此处写' in text} future={'平衡导航食品运输小车' in text}")
