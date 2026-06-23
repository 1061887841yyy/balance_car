from zipfile import ZipFile

from docx import Document


p = "outputs/FINAL_满版图文版_报告.docx"
d = Document(p)
all_text = "\n".join(x.text for x in d.paragraphs)
all_text += "\n" + "\n".join(cell.text for table in d.tables for row in table.rows for cell in row.cells)
with ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8")

print("title", "食品供应链双轮运输机器人设计与实现" in all_text)
print("group05", "第05组" in all_text)
print("members", all(name in all_text for name in ["陈俊杰", "冯超贤", "徐浤彬"]))
print("ids", all(sid in all_text for sid in ["2024433080104", "2024433080110", "2024433080138"]))
print("old_self_balance", "多传感器自平衡小车" in all_text or "自平衡小车设计与实现" in all_text)
print("new_terms", all(s in all_text for s in ["食品供应链", "双轮运输机器人", "蓝牙APP", "YOLO", "RFID", "保鲜"]))
print("qmarks", all_text.count("?"))
print("placeholder", "此处写" in all_text)
print("page_breaks", xml.count('w:type="page"'))
print("anchors", xml.count("<wp:anchor"))
print("wrapTopAndBottom", xml.count("<wp:wrapTopAndBottom"))
