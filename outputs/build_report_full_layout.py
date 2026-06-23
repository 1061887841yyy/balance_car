import importlib.util
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "source_template.docx"
IMG = ROOT / "outputs" / "图片材料"
ASSETS = ROOT / "outputs" / "report_full_assets"
OUT = ROOT / "outputs" / "FINAL_满版图文版_报告.docx"
OUT_ALIAS = ROOT / "outputs" / "FINAL_图文增强_8页正文_报告.docx"
OUT_ALIAS2 = ROOT / "outputs" / "FINAL_紧凑图文版_报告.docx"

spec = importlib.util.spec_from_file_location("build_report", ROOT / "outputs" / "build_report.py")
build_report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build_report)
TITLE = build_report.TITLE
PAGES = build_report.PAGES


def load_font(size):
    for path in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"]:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def crop_cover(path, size):
    im = Image.open(path).convert("RGB")
    return ImageOps.fit(im, size, method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))


def make_plate(items, out_name, title, cell=(420, 250), label_h=34):
    ASSETS.mkdir(exist_ok=True)
    font = load_font(22)
    title_font = load_font(28)
    pad = 18
    title_h = 42
    w = len(items) * cell[0] + (len(items) + 1) * pad
    h = title_h + cell[1] + label_h + pad * 2
    sheet = Image.new("RGB", (w, h), (250, 250, 250))
    draw = ImageDraw.Draw(sheet)
    draw.text((pad, 10), title, fill=(30, 30, 30), font=title_font)
    draw.line((pad, title_h, w - pad, title_h), fill=(170, 170, 170), width=2)
    for i, (file_name, label) in enumerate(items):
        x = pad + i * (cell[0] + pad)
        y = title_h + pad
        draw.rectangle([x - 1, y - 1, x + cell[0] + 1, y + cell[1] + label_h + 1], outline=(185, 185, 185), width=2)
        im = crop_cover(IMG / file_name, cell)
        sheet.paste(im, (x, y))
        draw.text((x + cell[0] / 2, y + cell[1] + 7), label, fill=(20, 20, 20), font=font, anchor="ma")
    out = ASSETS / out_name
    sheet.save(out, quality=92)
    return out


hardware_plate = make_plate(
    [
        ("主控stm32板.jpg", "STM32主控板"),
        ("mpu6050.jpg", "MPU6050姿态模块"),
        ("自研扩展版.png", "自研扩展板"),
    ],
    "plate_hardware.jpg",
    "硬件实物与扩展板",
)
sensor_plate = make_plate(
    [
        ("温湿度传感器.jpg", "DHT11温湿度"),
        ("柔性传感器.jpg", "柔性压力传感器"),
        ("oled.jpg", "OLED显示屏"),
    ],
    "plate_sensor_display.jpg",
    "温湿度、压力与显示模块",
)
design_plate = make_plate(
    [
        ("扩展版原理图1.png", "扩展板原理图"),
        ("模块支撑板.png", "模块支撑板"),
    ],
    "plate_design_future.jpg",
    "结构设计资料",
    cell=(560, 260),
)


def set_run_font(run, east="宋体", size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def format_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, east="宋体", bold=False, first_line=True, line_pt=24, before=0, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(24) if first_line else None
    for r in p.runs:
        set_run_font(r, east, size, bold)


def remove_paragraph(p):
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def add_after(paragraph, role, text=""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = paragraph._parent.paragraphs[[x._p for x in paragraph._parent.paragraphs].index(new_p)]
    if text:
        p.add_run(text)
    if role == "title":
        format_para(p, WD_ALIGN_PARAGRAPH.CENTER, 16, "黑体", True, False, 24, 0, 4)
    elif role == "heading":
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, 14, "楷体", True, False, 24, 8, 2)
    elif role == "body":
        format_para(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 12, "宋体", False, True, 24, 0, 0)
    elif role == "caption":
        format_para(p, WD_ALIGN_PARAGRAPH.CENTER, 10.5, "宋体", False, False, 16, 0, 4)
    elif role == "blank":
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, 12, "宋体", False, False, 12, 0, 0)
    return p


def add_image_after(paragraph, image_path, caption, width_cm=12.0):
    p_img = add_after(paragraph, "blank")
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.line_spacing = Pt(12)
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(0)
    p_img.add_run().add_picture(str(image_path), width=Cm(width_cm))
    p_cap = add_after(p_img, "caption", caption)
    return p_cap


def condensed_paragraphs(page):
    paras = []
    for text in page["paragraphs"]:
        paras.append(text)
    return paras


doc = Document(SOURCE)

# Remove only the original report-body placeholders. The front matter remains untouched.
for idx in range(62, 48, -1):
    remove_paragraph(doc.paragraphs[idx])

anchor = doc.paragraphs[48]
anchor.clear()
anchor.add_run(TITLE)
format_para(anchor, WD_ALIGN_PARAGRAPH.CENTER, 16, "黑体", True, False, 24, 0, 4)
current = anchor

image_inserted = {"hardware": False, "sensor": False, "display": False, "future": False}

for page in PAGES:
    heading = page["heading"].replace("（续）", "")
    # Avoid repeated visible headings caused by the old forced-page layout.
    if current.text.strip() != heading:
        current = add_after(current, "heading", heading)
    for para in condensed_paragraphs(page):
        current = add_after(current, "body", para)

        if (not image_inserted["hardware"]) and "硬件连接方面" in para:
            current = add_image_after(current, hardware_plate, "图1  主控、姿态传感器与扩展板集中展示", 12.0)
            note = "    图1将主控板、MPU6050姿态模块和自研扩展板放在同一版面中，便于对照硬件分工：STM32负责控制与数据处理，MPU6050负责姿态角测量，扩展板用于后续模块化连接和结构整理。"
            current = add_after(current, "body", note)
            image_inserted["hardware"] = True

        if (not image_inserted["sensor"]) and "DHT11驱动的设计重点" in para:
            current = add_image_after(current, sensor_plate, "图2  温湿度、压力与OLED显示模块集中展示", 12.0)
            note = "    图2对应本项目新增的测量与显示模块。DHT11用于环境温湿度采集，柔性压力传感器配合RFP602进行压力/重量估算，OLED作为现场数据显示窗口，三者共同补充了平衡车原有的姿态和速度测量链路。"
            current = add_after(current, "body", note)
            image_inserted["sensor"] = True

        if (not image_inserted["display"]) and "作品上电后" in para:
            current = add_image_after(current, design_plate, "图3  扩展板原理图与模块支撑板设计资料", 12.0)
            note = "    图3展示了扩展板原理图和模块支撑板设计资料。原理图用于说明电机驱动、传感器接口和电源输入等电路连接关系，支撑板则用于固定温湿度、压力、OLED等扩展模块，使硬件布局更加清晰，也便于后续维护和继续扩展。"
            current = add_after(current, "body", note)
            image_inserted["display"] = True

        if (not image_inserted["future"]) and "后续改进可以从控制性能" in para:
            future = "    在作品构想阶段，团队还曾规划过“平衡导航食品运输小车”方案：在自平衡底盘上增加雷达或循迹模块，使小车能够沿指定路径完成食品运输与供应。由于展示时间、调试风险和硬件稳定性限制，最终没有把该方案作为本次展示重点，但它为后续扩展提供了明确方向。未来可在现有姿态控制和传感器显示基础上，继续加入路径识别、避障、载物平台和通信上传功能。"
            current = add_after(current, "body", future)
            current = add_image_after(current, IMG / "平衡小车废案.jpg", "图4  废案构想：平衡导航食品运输小车", 6.8)
            image_inserted["future"] = True

# A final explicit page break is not needed; Word can flow the report naturally.
for p in doc.paragraphs:
    if not p.text.strip() and not p.runs:
        format_para(p, WD_ALIGN_PARAGRAPH.LEFT, 12, "宋体", False, False, 12, 0, 0)

doc.save(OUT)
for alias in [OUT_ALIAS, OUT_ALIAS2]:
    try:
        doc.save(alias)
    except PermissionError:
        print(f"skip_alias_locked={alias}")

with ZipFile(OUT) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")
text = "\n".join(p.text for p in doc.paragraphs)
print(f"saved={OUT}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"page_breaks={xml.count('w:type=\"page\"')} sectPr={xml.count('<w:sectPr')}")
print(f"qmarks={text.count('?')} placeholder={'此处写' in text} future={'平衡导航食品运输小车' in text}")
