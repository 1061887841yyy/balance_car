from pathlib import Path
from zipfile import ZipFile

from docx import Document


path = Path("outputs/report_filled_fixed.docx")
doc = Document(path)
text = "\n".join(p.text for p in doc.paragraphs)

with ZipFile(path) as zf:
    xml = zf.read("word/document.xml").decode("utf-8")

print(f"file={path.resolve()}")
print(f"paragraphs={len(doc.paragraphs)}")
print(f"tables={len(doc.tables)}")
print(f"xml_tables={xml.count('<w:tbl>')}")
print(f"sectPr={xml.count('<w:sectPr')}")
print(f"page_breaks={xml.count('w:type=\"page\"')}")
print(f"question_runs={xml.count('????')}")
print(f"has_placeholder={'此处写' in text}")
for i in [48, 50, 51, 56, 61, 66, 71, 76, 81, 86, 91, 96]:
    if i < len(doc.paragraphs):
        print(f"{i}: {doc.paragraphs[i].text[:80]}")
